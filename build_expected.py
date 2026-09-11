#!/usr/bin/env python3
import hashlib, json
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path('/Users/bp/cvent-one-shot')
rr=ROOT/'data/current/input.xlsx'; manifest_path=ROOT/'scope/intake-emerald-scope.json'
manifest=json.loads(manifest_path.read_text())
actual=hashlib.sha256((ROOT/'scope/intake-emerald.xlsx').read_bytes()).hexdigest()
if actual != manifest['sourceSha256']: raise SystemExit('scope hash mismatch')
confirmed={e['id']:e for e in manifest['entries'] if e['status']=='confirmed'}
wbv=load_workbook(rr,data_only=True); wbf=load_workbook(rr,data_only=False)
def fld(sid,value,source=None):
    if sid not in confirmed: raise ValueError(sid)
    d={'scopeId':sid,'value':value}
    if source:d['source']=source
    return d
def iso(v):
    return v.isoformat() if isinstance(v,(date,datetime)) else v
# Event basics
ws=wbv['Event Details']
location=str(ws['B10'].value).strip(); venue,city,state=[x.strip() for x in location.rsplit(',',2)]
event_fields={
 'venueName':fld('scope-007',venue,'Event Details!B10'),
 'city':fld('scope-008',city,'Event Details!B10'),
 'state':fld('scope-009',state,'Event Details!B10'),
 'timeZone':fld('scope-010',str(ws['B11'].value).strip(),'Event Details!B11'),
 'showHours':fld('scope-014',str(ws['B19'].value).strip(),'Event Details!B19'),
 'registrationDeadline':fld('scope-015','2026-11-15','NEW Reg Types & Pricing!U4'),
 'registrationGoal':fld('scope-016',ws['B22'].value,'Event Details!B22'),
}
# Theme, footer, body
links=wbv['Helpful & Social Media Links']
def yn(r): return str(links.cell(r,2).value or '').strip().lower()=='yes'
def link_field(sid,r): return fld(sid,{'visible':yn(r),'url':links.cell(r,3).value},f'Helpful & Social Media Links!A{r}:C{r}')
# Reg types and admission mappings
rpv=wbv['NEW Reg Types & Pricing']; rpf=wbf['NEW Reg Types & Pricing']
reg_by_code={}; admission={}
for r in range(5,28):
    code=rpv.cell(r,2).value; regname=rpv.cell(r,3).value; itemcode=rpv.cell(r,5).value; item=rpv.cell(r,6).value
    if not (code and regname): continue
    code=str(code).strip(); regname=str(regname).strip(); item=str(item).strip() if item else None
    rec=reg_by_code.setdefault(code,{'sourceIdentifier':regname,'fields':{'code':fld('scope-037',code,f'NEW Reg Types & Pricing!B{r}')}})
    if item:
        a=admission.setdefault(item,{'sourceCode':str(itemcode).strip() if itemcode else None,'fields':{'visibleToRegistrationTypeCodes':fld('scope-040',[],f'NEW Reg Types & Pricing!B{r}:F{r}')}})
        if code not in a['fields']['visibleToRegistrationTypeCodes']['value']: a['fields']['visibleToRegistrationTypeCodes']['value'].append(code)
        desc=rpv.cell(r,10).value
        if desc and 'description' not in a['fields']: a['fields']['description']=fld('scope-039',str(desc).strip(),f'NEW Reg Types & Pricing!J{r}')
# pricing records
pricing=[]
for r in range(5,28):
    code=rpv.cell(r,2).value; item=rpv.cell(r,6).value
    if not(code and item):continue
    pricing.append({'registrationTypeCode':str(code).strip(),'admissionItem':str(item).strip(),'fields':{
      'advancedPrice':fld('scope-044',rpv.cell(r,20).value,f'NEW Reg Types & Pricing!T{r}'),
      'onsitePrice':fld('scope-045',rpv.cell(r,21).value,f'NEW Reg Types & Pricing!U{r}')
    }})
# discounts
wsd=wbv['Discount Code Template']; discounts=[]
for r in range(7,wsd.max_row+1):
    name,code,method,amount=wsd.cell(r,1).value,wsd.cell(r,2).value,wsd.cell(r,4).value,wsd.cell(r,5).value
    if not(name and code and method and amount is not None):continue
    fields={
      'name':fld('scope-047',str(name).strip(),f'Discount Code Template!A{r}'),
      'code':fld('scope-048',str(code).strip(),f'Discount Code Template!B{r}'),
      'method':fld('scope-049',str(method).strip(),f'Discount Code Template!D{r}'),
      'value':fld('scope-050',amount,f'Discount Code Template!E{r}'),
      'effectiveFromTo':fld('scope-051',{'from':iso(wsd.cell(r,6).value),'to':iso(wsd.cell(r,7).value)},f'Discount Code Template!F{r}:G{r}'),
      'capacity':fld('scope-052',wsd.cell(r,8).value,f'Discount Code Template!H{r}'),
      'stackable':fld('scope-053',str(wsd.cell(r,9).value).strip(),f'Discount Code Template!I{r}'),
      'usableBy':fld('scope-054',str(wsd.cell(r,10).value).strip(),f'Discount Code Template!J{r}'),
      'countGuestsTowardCapacity':fld('scope-055',str(wsd.cell(r,11).value).strip(),f'Discount Code Template!K{r}'),
      'active':fld('scope-056',str(wsd.cell(r,12).value).strip(),f'Discount Code Template!L{r}'),
      'applicableAdmissionItemCodes':fld('scope-057',[x.strip() for x in str(wsd.cell(r,14).value or '').split(',') if x.strip()],f'Discount Code Template!N{r}')
    }
    discounts.append({'fields':fields})
# questions: headers with display text D and one or more control columns; append option rows and standalone D text until next header.
wsq=wbv['Show Questions']; questions=[]; current=None
for r in range(5,wsq.max_row+1):
    a,b,c,d,e,f,g,h,i=[wsq.cell(r,x).value for x in range(1,10)]
    is_header=bool(d is not None and (a is not None or b is not None or c is not None or g is not None or h is not None or i is not None))
    if is_header:
        fields={}
        if a is not None: fields['pagePlacement']=fld('scope-058',str(a).strip(),f'Show Questions!A{r}')
        if c is not None: fields['companyOrIndividual']=fld('scope-059',str(c).strip(),f'Show Questions!C{r}')
        fields['displayedText']=fld('scope-060',str(d).strip(),f'Show Questions!D{r}')
        if g is not None: fields['appearance']=fld('scope-061',str(g).strip(),f'Show Questions!G{r}')
        if h is not None: fields['required']=fld('scope-063',str(h).strip(),f'Show Questions!H{r}')
        if i is not None: fields['registrationTypeVisibility']=fld('scope-064',str(i).strip(),f'Show Questions!I{r}')
        current={'sourceIdentifier':str(b).strip() if b is not None else None,'sourceRow':r,'fields':fields}
        questions.append(current)
    elif current:
        if e is not None or f is not None:
            opts=current['fields'].setdefault('answerOptions',fld('scope-062',[],f'Show Questions!E{current["sourceRow"]}:F{r}'))['value']
            opts.append({'code':str(e).strip() if e is not None else None,'text':str(f).strip() if f is not None else None})
        elif d is not None:
            current['fields']['displayedText']['value'] += '\n\n'+str(d).strip()
# expected output
out={
 'schemaVersion':1,'runMode':'mock','authority':'Forge Intake',
 'scopeManifest':{'path':str(manifest_path),'sourceWorkbookSha256':actual,'verified':True},
 'rr':{'path':str(rr),'sha256':hashlib.sha256(rr.read_bytes()).hexdigest()},
 'targetGuardrail':{'authorizedEventName':'(C+D) Medtrade Testing Clone 2','eventKey':'e712e34c-6117-4d13-bf4c-8ed54cf2b495','preserveNameCodeKeyUrlAndUnpublished':True},
 'domains':{
  'event_basics':{'fields':event_fields},
  'theme_branding':{'fields':{'eventTheme':fld('scope-017',str(ws['B14'].value).strip(),'Event Details!B14'),'brandPrimaryColors':fld('scope-019',[x.split('#',1)[1].strip() for x in str(ws['B15'].value).splitlines() if '#' in x],'Event Details!B15')}},
  'header_footer_body':{'fields':{
   'attendeeShowHoursLink':link_field('scope-024',3),'attendeeShowPolicyLink':link_field('scope-025',4),'attendeeEmeraldPrivacyPolicyLink':link_field('scope-026',5),'attendeeBrowseSessionsLink':link_field('scope-027',7),'attendeeReviewPricingLink':link_field('scope-028',8),'attendeeFaqLink':link_field('scope-030',9),'attendeeContactUs':link_field('scope-031',10),
   'exhibitorShowHoursLink':link_field('scope-024',15),'exhibitorShowPolicyLink':link_field('scope-025',16),'exhibitorEmeraldPrivacyPolicyLink':link_field('scope-026',17),'exhibitorBrowseSessionsLink':link_field('scope-027',19),'exhibitorReviewPricingLink':link_field('scope-028',20),'exhibitorFaqLink':link_field('scope-030',21),'exhibitorResourceCenter':link_field('scope-032',22),'exhibitorContactUs':link_field('scope-031',23),
   'countdownTimer':fld('scope-033',{'visible':yn(28),'text':links['B29'].value},'Helpful & Social Media Links!B28:B29'),
   'socialFollowBar':fld('scope-034',[{'type':links.cell(r,1).value,'url':links.cell(r,3).value} for r in range(32,37) if yn(r)],'Helpful & Social Media Links!A32:C36'),
   'alreadyRegisteredLink':fld('scope-035',str(ws['B35'].value).strip(),'Event Details!B35')
  }},
  'registration_paths':{'fields':{}},
  'registration_types':{'objects':list(reg_by_code.values())},
  'admission_items':{'objects':[{'sourceIdentifier':k,**v} for k,v in admission.items()]},
  'pricing_fees':{'fields':{'priceTierDateRanges':fld('scope-041',[{'name':'Advanced','from':'2026-06-01','to':'2026-11-12'},{'name':'Onsite','from':'2026-11-13','to':'2026-11-15'}],'NEW Reg Types & Pricing!T4:U4'),'processingFeeNote':fld('scope-046','A 3% processing fee will be added to the total balance at the end of registration for all shows.','NEW Reg Types & Pricing!B1')},'objects':pricing},
  'discounts':{'objects':discounts},
  'registration_questions':{'objects':questions},
  'terms_policies':{'fields':{'showPolicyUrl':fld('scope-066',links['C4'].value,'Helpful & Social Media Links!C4')},'objects':[q for q in questions if q.get('sourceIdentifier') in ('POLICY','EMPP')]},
 },
 'blockedRequests':[
  {'scopeId':'scope-005','request':str(ws['B9'].value).strip(),'reason':'Protected event identity must be preserved; mock RR cannot rename or select the target.'},
  {'scopeId':'scope-038','request':'REG TYPE NAME column','reason':'Current RR header maps to unconfirmed Reg Type Name (scope-078), not the confirmed Registration Pass Description field; fail-closed.'},
  {'scopeIds':['scope-058','scope-059','scope-060','scope-061','scope-062','scope-063','scope-064'],'request':'Create any missing registration question','reason':'Creation requires an unconfirmed internal name (scope-103); no internal name will be invented or changed.'}
 ],
 'excludedRequests':[
  {'scopeId':'scope-071','request':'Hotel Info links','status':'unconfirmed'},
  {'scopeIds':['scope-078','scope-080','scope-081','scope-082','scope-083','scope-084','scope-085'],'request':'Registration type name, path/status/schedule/capacity/guest controls','status':'unconfirmed'},
  {'scopeIds':['scope-086','scope-087','scope-088','scope-089','scope-090','scope-091'],'request':'Admission item identity/status/schedule/capacity/fee controls','status':'unconfirmed'},
  {'scopeIds':['scope-097','scope-098'],'request':'Discount sessions and optional-item applicability','status':'unconfirmed'},
  {'scopeIds':['scope-103','scope-104','scope-105'],'request':'Question internal names, determines-registration-type logic, triggers/follow-ups','status':'unconfirmed'},
  {'request':'Integrations, communications, approvals, sessions, optional items, group/volume discounts, badge/ticket layouts, onsite operations/scanners, evaluations and reference/global definitions','status':'absent or deferred'}
 ]
}
# Validate every fields leaf has an exact confirmed scopeId.
def walk(v,path='root'):
    if isinstance(v,dict):
        if 'fields' in v:
            for k,x in v['fields'].items():
                if not isinstance(x,dict) or x.get('scopeId') not in confirmed: raise ValueError(f'{path}.fields.{k}')
        for k,x in v.items(): walk(x,path+'.'+k)
    elif isinstance(v,list):
        for i,x in enumerate(v): walk(x,f'{path}[{i}]')
walk(out['domains'])
path=ROOT/'data/current/expected-domains.json'; path.write_text(json.dumps(out,indent=2,ensure_ascii=False))
counts={k:sum(len(v.get('fields',{}))+sum(len(o.get('fields',{})) for o in v.get('objects',[])) for _ in [0]) for k,v in out['domains'].items()}
print(json.dumps({'output':str(path),'domainFieldInstances':counts,'discounts':len(discounts),'questions':len(questions),'registrationTypes':len(reg_by_code),'admissionItems':len(admission)},indent=2))
