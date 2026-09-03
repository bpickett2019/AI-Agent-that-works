import json
import tempfile
import unittest
from pathlib import Path

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

from workbook_ops import sheet, update


class WorkbookEditingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        (self.root / "activity.log").touch()
        (self.root / "state.json").write_text(json.dumps({
            "status": "draft", "current_action": "ready", "rr_file": "rr.xlsx",
            "authorized_event_name": "Authorized test", "authorized_event_id": "event-1",
            "authorized_event_key": "event-1",
        }))

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "RR"
        worksheet["A1"] = "Original"
        worksheet["B1"] = 12
        worksheet["C1"] = "Merged"
        worksheet.merge_cells("C1:D1")
        workbook.save(self.root / "input.xlsx")
        workbook.close()

    def tearDown(self):
        self.temp.cleanup()

    def version(self):
        return str((self.root / "input.xlsx").stat().st_mtime_ns)

    def test_batch_edit_saves_backup_preserves_types_and_resets_state(self):
        (self.root / "expected-domains.json").write_text("{}")
        result = update(self.root, {
            "version": self.version(),
            "changes": [
                {"sheet": "RR", "row": 1, "column": 1, "value": "Changed"},
                {"sheet": "RR", "row": 1, "column": 2, "value": "42"},
                {"sheet": "RR", "row": 2, "column": 1, "value": "=B1*2"},
            ],
        }, is_running=False)
        self.assertEqual(result["saved"], 3)
        self.assertEqual(result["version"], self.version())
        self.assertTrue((self.root / "workbook-backups" / result["backup"]).exists())
        self.assertFalse((self.root / "expected-domains.json").exists())

        workbook = load_workbook(self.root / "input.xlsx", data_only=False)
        try:
            self.assertEqual(workbook["RR"]["A1"].value, "Changed")
            self.assertEqual(workbook["RR"]["B1"].value, 42)
            self.assertIsInstance(workbook["RR"]["B1"].value, int)
            self.assertEqual(workbook["RR"]["A2"].value, "=B1*2")
        finally:
            workbook.close()
        state = json.loads((self.root / "state.json").read_text())
        self.assertEqual(state["status"], "draft")
        self.assertIn("edited", state["current_action"])
        self.assertEqual(state["authorized_event_id"], "event-1")

    def test_rejects_stale_version_merged_cell_and_running_agent(self):
        with self.assertRaises(HTTPException) as stale:
            update(self.root, {"version": "stale", "changes": [
                {"sheet": "RR", "row": 1, "column": 1, "value": "x"}
            ]}, is_running=False)
        self.assertEqual(stale.exception.status_code, 409)

        with self.assertRaises(HTTPException) as merged:
            update(self.root, {"version": self.version(), "changes": [
                {"sheet": "RR", "row": 1, "column": 4, "value": "x"}
            ]}, is_running=False)
        self.assertEqual(merged.exception.status_code, 409)

        with self.assertRaises(HTTPException) as active:
            update(self.root, {"version": self.version(), "changes": [
                {"sheet": "RR", "row": 1, "column": 1, "value": "x"}
            ]}, is_running=True)
        self.assertEqual(active.exception.status_code, 409)

    def test_sheet_marks_merged_follower_non_editable(self):
        body = sheet(self.root, "RR", 1, 10)
        self.assertTrue(body["editable"][0][2])
        self.assertFalse(body["editable"][0][3])
        self.assertIsInstance(body["version"], str)


if __name__ == "__main__":
    unittest.main()
