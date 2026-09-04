import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]


class RRCompilerTests(unittest.TestCase):
    def compile_with_registration_sheet(self, sheet_name):
        with tempfile.TemporaryDirectory() as temporary:
            job_dir = Path(temporary)
            workbook = Workbook()
            workbook.remove(workbook.active)

            event = workbook.create_sheet("Event Details")
            event["A1"] = "Event Location"
            event["B1"] = "Test Venue, Test City, TS"

            links = workbook.create_sheet("Helpful & Social Media Links")
            links["B28"] = "No"

            registration = workbook.create_sheet(sheet_name)
            registration["B1"] = "No processing fee"
            registration["T4"] = "Advance"
            registration["U4"] = "Onsite 01/01/2030 - 01/02/2030"
            registration["B5"] = "REG-1"
            registration["C5"] = "Test registration type"
            registration["D5"] = "ACTIVATE"
            registration["E5"] = "ITEM-1"
            registration["F5"] = "Test admission item"
            registration["T5"] = 10
            registration["U5"] = 20

            workbook.create_sheet("Discount Code Template")
            workbook.create_sheet("Show Questions")
            workbook.save(job_dir / "input.xlsx")
            workbook.close()

            environment = os.environ.copy()
            environment["CVENT_JOB_DIR"] = str(job_dir)
            completed = subprocess.run(
                [sys.executable, str(ROOT / "rr_compiler.py")],
                cwd=ROOT,
                env=environment,
                text=True,
                capture_output=True,
                check=True,
            )
            result = json.loads(completed.stdout)
            expected = json.loads((job_dir / "expected-domains.json").read_text())
            return result, expected

    def test_accepts_supported_registration_sheet_names_and_records_exact_source(self):
        for sheet_name in ("NEW Reg Types & Pricing", "Reg Types & Pricing"):
            with self.subTest(sheet_name=sheet_name):
                result, expected = self.compile_with_registration_sheet(sheet_name)
                self.assertTrue(result["ok"])
                item = expected["domains"]["registration_types"]["items"][0]
                self.assertEqual(item["fields"]["code"]["source"], f"{sheet_name}!B5")
                pricing = expected["domains"]["pricing_fees"]
                self.assertEqual(pricing["fields"]["tier_date_ranges"]["source"], f"{sheet_name}!T4:U4")


if __name__ == "__main__":
    unittest.main()
