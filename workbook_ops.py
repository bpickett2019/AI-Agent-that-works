"""Job-scoped Excel preview and safe atomic editing."""
from __future__ import annotations

import os
import shutil
from datetime import date, datetime, time as datetime_time
from pathlib import Path
from threading import Lock

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

from job_runner import atomic_json, now, read_json

_LOCKS: dict[str, Lock] = {}


def lock_for(directory: Path) -> Lock:
    return _LOCKS.setdefault(str(directory.resolve()), Lock())


def info(directory: Path) -> dict:
    path = directory / "input.xlsx"
    if not path.exists():
        raise HTTPException(404, "No RR workbook uploaded")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = [{"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column}
                  for sheet in workbook.worksheets]
    finally:
        workbook.close()
    state = read_json(directory / "state.json", {})
    return {"file": state.get("rr_file") or path.name, "version": str(path.stat().st_mtime_ns), "sheets": sheets}


def sheet(directory: Path, name: str, start: int = 1, limit: int = 80) -> dict:
    path = directory / "input.xlsx"
    if not path.exists():
        raise HTTPException(404, "No RR workbook uploaded")
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if name not in workbook.sheetnames:
            raise HTTPException(404, "Worksheet not found")
        worksheet = workbook[name]
        start = max(1, start)
        limit = max(10, min(limit, 150))
        end = min(worksheet.max_row, start + limit - 1)
        width = min(worksheet.max_column, 60)

        def value(item):
            if item is None:
                return ""
            if hasattr(item, "isoformat"):
                return item.isoformat()
            return str(item)

        rows = [[value(worksheet.cell(row, column).value) for column in range(1, width + 1)]
                for row in range(start, end + 1)]
        editable = [[not worksheet.protection.sheet and not isinstance(worksheet.cell(row, column), MergedCell)
                     for column in range(1, width + 1)] for row in range(start, end + 1)]
        return {
            "name": name, "version": str(path.stat().st_mtime_ns), "start": start, "end": end,
            "total_rows": worksheet.max_row, "total_columns": worksheet.max_column,
            "columns": [get_column_letter(column) for column in range(1, width + 1)],
            "rows": rows, "editable": editable, "protected": bool(worksheet.protection.sheet),
        }
    finally:
        workbook.close()


def edited_cell_value(cell, text):
    if text == "":
        return None
    if text.startswith("="):
        return text
    current = cell.value
    try:
        if isinstance(current, bool):
            lowered = text.strip().lower()
            if lowered not in ("true", "false"):
                raise ValueError
            return lowered == "true"
        if isinstance(current, int) and not isinstance(current, bool):
            return int(text)
        if isinstance(current, float):
            return float(text)
        if isinstance(current, datetime):
            return datetime.fromisoformat(text)
        if isinstance(current, date):
            return date.fromisoformat(text)
        if isinstance(current, datetime_time):
            return datetime_time.fromisoformat(text)
    except ValueError as exc:
        raise HTTPException(400, f"Value {text!r} is invalid for {cell.coordinate}") from exc
    return text


def reset_after_edit(directory: Path, filename: str) -> None:
    for name in (
        "benchmark-results.json", "build-checklist.json", "domain-results.json", "expected-domains.json",
        "input.inspection.json", "input.inspection-summary.json", "review-required.json", "rr-checklist.json",
        "rr-execution-checklist.json", "rr-focus.json", "rr-focused.txt", "rr-inspection.json",
        "rr-inspection-summary.json", "rr-question-checklist.json",
    ):
        try:
            (directory / name).unlink()
        except FileNotFoundError:
            pass
    state = read_json(directory / "state.json", {})
    state.update({
        "status": "draft", "current_stage": "upload",
        "current_action": "RR workbook edited — ready to reread requirements",
        "completed": [], "review_required": [], "rr_file": filename,
        "pi_pid": None, "pi_session": None, "updated_at": now(),
    })
    atomic_json(directory / "state.json", state)
    atomic_json(directory / "final-report.json", {
        "status": "INCOMPLETE", "unresolved_items": ["Workbook changed; build must reread requirements"],
        "real_reads": [], "real_writes": [],
        "guardrails": {"published": 0, "emails_sent": 0, "deletes": 0, "global_mutations": 0},
        "updated_at": now(),
    })


def update(directory: Path, payload: dict, is_running: bool) -> dict:
    path = directory / "input.xlsx"
    if not path.exists():
        raise HTTPException(404, "No RR workbook uploaded")
    if is_running:
        raise HTTPException(409, "Stop CVENT Agent before editing the RR workbook")
    changes = payload.get("changes")
    version = payload.get("version")
    if not isinstance(changes, list) or not changes or len(changes) > 2000:
        raise HTTPException(400, "Submit between 1 and 2000 cell changes")
    with lock_for(directory):
        if str(version) != str(path.stat().st_mtime_ns):
            raise HTTPException(409, "The workbook changed; reload it before saving")
        workbook = load_workbook(path, data_only=False)
        temporary = path.with_name("input.editing.xlsx")
        try:
            normalized = []
            seen = set()
            for change in changes:
                if not isinstance(change, dict):
                    raise HTTPException(400, "Each cell change must be an object")
                sheet_name = change.get("sheet")
                row = change.get("row")
                column = change.get("column")
                value = change.get("value")
                if not isinstance(sheet_name, str) or sheet_name not in workbook.sheetnames:
                    raise HTTPException(400, "Unknown worksheet")
                if not isinstance(row, int) or not 1 <= row <= 1048576 or not isinstance(column, int) or not 1 <= column <= 16384:
                    raise HTTPException(400, "Invalid cell coordinates")
                if not isinstance(value, str) or len(value) > 32767:
                    raise HTTPException(400, "Cell values must be text no longer than 32,767 characters")
                key = (sheet_name, row, column)
                if key in seen:
                    raise HTTPException(400, "Duplicate cell change")
                seen.add(key)
                worksheet = workbook[sheet_name]
                cell = worksheet.cell(row, column)
                if worksheet.protection.sheet:
                    raise HTTPException(409, f"Worksheet {sheet_name} is protected")
                if isinstance(cell, MergedCell):
                    raise HTTPException(409, f"{sheet_name}!{cell.coordinate} is a non-editable merged cell")
                normalized.append((cell, edited_cell_value(cell, value)))
            backups = directory / "workbook-backups"
            backups.mkdir(parents=True, exist_ok=True)
            stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
            backup = backups / f"input-{stamp}.xlsx"
            shutil.copy2(path, backup)
            for cell, value in normalized:
                cell.value = value
            workbook.save(temporary)
            os.replace(temporary, path)
        finally:
            workbook.close()
            try:
                temporary.unlink()
            except FileNotFoundError:
                pass
        filename = read_json(directory / "state.json", {}).get("rr_file") or "input.xlsx"
        reset_after_edit(directory, filename)
        return {"ok": True, "saved": len(normalized), "version": str(path.stat().st_mtime_ns), "backup": backup.name}
