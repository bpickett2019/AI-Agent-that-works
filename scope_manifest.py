#!/usr/bin/env python3
"""Compile and verify the authoritative Forge Intake automation scope."""
from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
SCOPE_WORKBOOK = ROOT / 'scope' / 'intake-emerald.xlsx'
SCOPE_MANIFEST = ROOT / 'scope' / 'intake-emerald-scope.json'
STATUS_HEADINGS = {
    'confirmed': 'confirmed',
    'unconfirmed': 'unconfirmed',
    'defered to post mvp': 'deferred',
    'deferred to post mvp': 'deferred',
}


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def clean(value):
    return value.strip() if isinstance(value, str) else value


def compile_scope(path: Path = SCOPE_WORKBOOK) -> dict:
    path = Path(path)
    wb = load_workbook(path, data_only=False, read_only=False)
    try:
        ws = wb.active
        status = None
        section = None
        entries = []
        for row_number in range(1, ws.max_row + 1):
            values = [clean(ws.cell(row_number, column).value) for column in range(1, 9)]
            populated = [value for value in values if value not in (None, '')]
            heading = str(values[0]).strip().lower() if values[0] not in (None, '') else ''
            if len(populated) == 1 and heading in STATUS_HEADINGS:
                status = STATUS_HEADINGS[heading]
                section = None
                continue
            if row_number == 3 or not status or values[3] in (None, ''):
                continue
            if values[0] not in (None, ''):
                section = values[0]
            entry = {
                'id': f'scope-{row_number:03d}',
                'row': row_number,
                'status': status,
                'section': section,
                'cventNavigation': values[1],
                'cventArea': values[2],
                'cventField': values[3],
                'rrTab': values[4],
                'rrField': values[5],
                'notes': values[6],
                'approval': values[7],
            }
            entries.append(entry)
        counts = {name: sum(entry['status'] == name for entry in entries) for name in ('confirmed', 'unconfirmed', 'deferred')}
        return {
            'schemaVersion': 1,
            'authority': 'Forge Intake',
            'sourceWorkbook': 'scope/intake-emerald.xlsx',
            'sourceWorksheet': ws.title,
            'sourceSha256': file_sha256(path),
            'policy': {
                'confirmed': 'automation may inspect and modify only these mapped Cvent fields',
                'unconfirmed': 'review and report only; no automated Cvent writes',
                'deferred': 'out of current MVP execution; do not inspect or modify',
                'absent': 'out of scope; do not inspect or modify',
            },
            'counts': counts,
            'confirmedIds': [entry['id'] for entry in entries if entry['status'] == 'confirmed'],
            'confirmedNavigation': sorted({entry['cventNavigation'] for entry in entries if entry['status'] == 'confirmed' and entry['cventNavigation']}),
            'entries': entries,
        }
    finally:
        wb.close()


def write_manifest(workbook: Path = SCOPE_WORKBOOK, output: Path = SCOPE_MANIFEST) -> dict:
    manifest = compile_scope(workbook)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix('.tmp')
    temporary.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + '\n')
    temporary.replace(output)
    return manifest


def load_manifest(workbook: Path = SCOPE_WORKBOOK, manifest_path: Path = SCOPE_MANIFEST) -> dict:
    manifest = json.loads(Path(manifest_path).read_text())
    actual_hash = file_sha256(Path(workbook))
    if manifest.get('sourceSha256') != actual_hash:
        raise RuntimeError('Automation scope workbook and manifest do not match; regenerate the scope manifest')
    if not manifest.get('confirmedIds') or not manifest.get('entries'):
        raise RuntimeError('Automation scope manifest is empty')
    return manifest


if __name__ == '__main__':
    source = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else SCOPE_WORKBOOK
    destination = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else SCOPE_MANIFEST
    result = write_manifest(source, destination)
    print(json.dumps({'ok': True, 'output': str(destination), 'sha256': result['sourceSha256'], 'counts': result['counts']}, indent=2))
