#!/usr/bin/env python3
"""Independent deterministic RR evidence validation and full execution planning."""
from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
from datetime import date, datetime, time, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

ROOT = Path(__file__).resolve().parent
JOB_DIR = Path(os.environ.get("CVENT_JOB_DIR", ROOT / "data/current"))
RR = JOB_DIR / "input.xlsx"
EXPECTED = JOB_DIR / "expected-domains.json"
VALIDATION = JOB_DIR / "rr-validation.json"
PLAN = JOB_DIR / "configuration-plan.json"
ORDER = [
    "event_settings", "registration_types", "admission_items", "optional_items", "pricing",
    "discounts_vouchers", "questions", "sessions", "registration_paths", "site_designer",
    "integrations", "communications", "badges_onsite", "associations", "final_qa",
]
PROCEDURES = {
    "event_settings": "configure_event_settings", "registration_types": "configure_registration_types",
    "admission_items": "configure_admission_items", "optional_items": "configure_optional_items",
    "pricing": "configure_pricing", "discounts_vouchers": "import_or_configure_discounts",
    "questions": "configure_questions_and_choices", "sessions": "configure_sessions",
    "registration_paths": "configure_registration_paths", "site_designer": "configure_site_designer",
    "integrations": "configure_integrations", "communications": "configure_event_communications",
    "badges_onsite": "configure_badges_and_onsite", "associations": "configure_event_associations",
    "final_qa": "verify_rr_against_cvent",
}


def clean(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def atomic_json(path: Path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, indent=2, ensure_ascii=False)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.chmod(temporary, 0o600)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def normalize(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def source_values(workbook, source):
    if not isinstance(source, str) or "!" not in source:
        raise ValueError("missing RR source")
    sheet_name, coordinate = source.rsplit("!", 1)
    sheet_name = sheet_name.strip("'")
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"source sheet is absent: {sheet_name}")
    ws = workbook[sheet_name]
    min_col, min_row, max_col, max_row = range_boundaries(coordinate)
    values = []
    for row in range(min_row, max_row + 1):
        for column in range(min_col, max_col + 1):
            value = clean(ws.cell(row, column).value)
            if value is not None:
                values.append(value)
    return ws, (min_col, min_row, max_col, max_row), values


def section_context(ws, bounds):
    min_col, min_row, max_col, _ = bounds
    labels = []
    for row in range(max(1, min_row - 8), min_row + 1):
        for column in range(1, min(ws.max_column, max_col + 2) + 1):
            if row == min_row and min_col <= column <= max_col:
                continue
            value = clean(ws.cell(row, column).value)
            if isinstance(value, str) and value not in labels:
                labels.append(value)
    return labels[-6:]


def desired_scalars(value):
    if isinstance(value, dict):
        return [item for key, child in value.items() if key not in {"source", "sourceRow"} for item in desired_scalars(child)]
    if isinstance(value, list):
        return [item for child in value for item in desired_scalars(child)]
    return [] if value is None else [value]


def independently_supported(desired, raw_values):
    raw = [normalize(value) for value in raw_values if normalize(value)]
    wanted = [normalize(value) for value in desired_scalars(desired) if normalize(value)]
    if not raw:
        return False
    if isinstance(desired, bool):
        truthy = raw[0] in {"yes", "y", "true", "1", "activate", "active", "required"}
        falsy = raw[0] in {"no", "n", "false", "0", "inactive"}
        return truthy if desired else falsy
    if all(item in raw for item in wanted):
        return True
    combined_raw = " ".join(raw)
    return bool(wanted) and all(item in combined_raw or combined_raw in item for item in wanted)


def evidence_nodes(value, path=()):
    if isinstance(value, dict):
        if isinstance(value.get("source"), str):
            desired = value.get("value") if "value" in value else {key: child for key, child in value.items() if key != "source"}
            yield path, value["source"], desired
            return
        for key, child in value.items():
            yield from evidence_nodes(child, path + (str(key),))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from evidence_nodes(child, path + (str(index),))


def main():
    expected = json.loads(EXPECTED.read_text(encoding="utf-8"))
    workbook = load_workbook(RR, data_only=True, read_only=False)
    items = []
    context_cache = {}
    try:
        for domain in ORDER:
            for path, source, desired in evidence_nodes(expected.get("domains", {}).get(domain, {}), (domain,)):
                item_id = hashlib.sha256(("/".join(path) + "|" + source).encode()).hexdigest()[:20]
                try:
                    ws, bounds, raw_values = source_values(workbook, source)
                    status = "VERIFIED" if independently_supported(desired, raw_values) else "AMBIGUOUS"
                    error = None
                    context_key = (ws.title, bounds[1])
                    if context_key not in context_cache:
                        context_cache[context_key] = section_context(ws, bounds)
                    context = context_cache[context_key]
                except Exception as exc:
                    raw_values, context, status, error = [], [], "NOT_SUPPORTED_BY_RR", str(exc)
                items.append({
                    "itemId": item_id, "domain": domain, "path": "/".join(path), "status": status,
                    "sourceEvidence": {"sheet": source.rsplit("!", 1)[0] if "!" in source else None,
                        "range": source.rsplit("!", 1)[1] if "!" in source else None,
                        "surroundingHeaderOrSection": context, "rawWorkbookValue": raw_values},
                    "interpretedCventValue": desired, **({"validationError": error} if error else {}),
                })
    finally:
        workbook.close()
    counts = {status: sum(item["status"] == status for item in items) for status in ("VERIFIED", "AMBIGUOUS", "NOT_SUPPORTED_BY_RR")}
    meta = {
        "schemaVersion": 1, "generatedAt": datetime.now(timezone.utc).isoformat(),
        "rrSha256": hashlib.sha256(RR.read_bytes()).hexdigest(),
        "expectationsSha256": hashlib.sha256(EXPECTED.read_bytes()).hexdigest(), "counts": counts,
    }
    atomic_json(VALIDATION, {**meta, "items": items})
    sections = []
    for domain in ORDER:
        domain_items = [item for item in items if item["domain"] == domain]
        sections.append({"order": len(sections) + 1, "domain": domain, "procedure": PROCEDURES[domain],
            "verifiedItemIds": [item["itemId"] for item in domain_items if item["status"] == "VERIFIED"],
            "heldItems": [{"itemId": item["itemId"], "status": item["status"]} for item in domain_items if item["status"] != "VERIFIED"]})
    atomic_json(PLAN, {**meta, "target": expected.get("target"), "mission": sections,
        "executionRule": "Execute VERIFIED items, hold only AMBIGUOUS/NOT_SUPPORTED_BY_RR items, and account for every item in final verification."})
    print(json.dumps({"ok": True, "validation": str(VALIDATION), "plan": str(PLAN), "counts": counts}, indent=2))


if __name__ == "__main__":
    main()
