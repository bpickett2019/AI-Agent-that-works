#!/usr/bin/env python3
"""Compile writable event-configuration requirements directly from the uploaded RR."""
from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
from datetime import date, datetime, time, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
JOB_DIR = Path(os.environ.get("CVENT_JOB_DIR", ROOT / "data/current"))
RR = JOB_DIR / "input.xlsx"
OUT = JOB_DIR / "expected-domains.json"
TARGET = {
    "name": os.environ.get("CVENT_AUTHORIZED_EVENT_NAME", "(C+D) Medtrade Testing Clone 2"),
    "eventId": os.environ.get("CVENT_AUTHORIZED_EVENT_ID", ""),
    "eventKey": os.environ.get("CVENT_AUTHORIZED_EVENT_KEY", "e712e34c-6117-4d13-bf4c-8ed54cf2b495"),
    "eventCode": os.environ.get("CVENT_AUTHORIZED_EVENT_CODE", ""),
    "mustRemainUnpublished": True,
}


def normalized(value):
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return value


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.replace("\xa0", " ").strip()
        return value or None
    return normalized(value)


def yes_no(value):
    text = str(value or "").strip().lower()
    if text in {"yes", "y", "true", "required", "activate"}:
        return True
    if text in {"no", "n", "false", "not needed"}:
        return False
    return clean(value)


def field(value, source, **extra):
    item = {"value": clean(value), "source": source}
    item.update(extra)
    return item


def atomic_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=path.name + ".", dir=path.parent)
    try:
        with os.fdopen(fd, "w") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def write_discount_import(path, source_ws, columns_by_field, records):
    workbook = Workbook(write_only=False)
    sheet = workbook.active
    sheet.title = "Discount Codes"
    schema = [
        ("Name", "name"), ("Discount Code", "code"), ("Discount Type", "type"), ("Method", "method"),
        ("Amount/Percentage", "amount_or_percentage"), ("Effective From", "effective_from"), ("Effective To", "effective_to"),
        ("Capacity", "capacity"), ("Stackable", "stackable"), ("Can be used by", "usable_by"),
        ("Counts guests", "count_guests"), ("Active", "active"), ("Internal Note", "internal_note"),
        ("Admission Items", "admission_items"), ("Tracks", "tracks"), ("Sessions", "sessions"),
    ]
    sheet.append([header for header, _ in schema])
    method_index = next(index for index, (_, name) in enumerate(schema) if name == "method")
    amount_index = next(index for index, (_, name) in enumerate(schema) if name == "amount_or_percentage")
    for record in records:
        row = record["sourceRow"]
        values = [source_ws.cell(row, columns_by_field[name]).value if columns_by_field.get(name) else None for _, name in schema]
        method = str(clean(values[method_index]) or "")
        if re.fullmatch(r"amount off", method, re.I):
            values[method_index] = "Subtract an amount"
        elif re.fullmatch(r"percent(?:age)? off", method, re.I):
            values[method_index] = "Subtract a percentage"
        amount_column = columns_by_field.get("amount_or_percentage")
        if "percentage" in str(values[method_index]).lower() and amount_column and "%" in source_ws.cell(row, amount_column).number_format:
            amount = values[amount_index]
            if isinstance(amount, (int, float)):
                values[amount_index] = amount * 100
        sheet.append(values)
    temporary = path.with_name(path.name + ".tmp.xlsx")
    workbook.save(temporary)
    workbook.close()
    os.chmod(temporary, 0o600)
    os.replace(temporary, path)
    return {"path": str(path), "records": len(records), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def value_map(ws):
    return {
        str(ws.cell(row, 1).value).strip(): clean(ws.cell(row, 2).value)
        for row in range(1, ws.max_row + 1)
        if clean(ws.cell(row, 1).value) is not None
    }


def normalized_label(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def worksheet_by_alias(workbook, purpose, *names):
    aliases = [normalized_label(name) for name in names]
    exact = [name for name in workbook.sheetnames if normalized_label(name) in aliases]
    if len(exact) == 1:
        return exact[0], workbook[exact[0]]
    scored = []
    for sheet_name in workbook.sheetnames:
        sheet_tokens = set(normalized_label(sheet_name).split())
        score = max((len(sheet_tokens & set(alias.split())) / max(1, len(sheet_tokens | set(alias.split()))) for alias in aliases), default=0)
        if score >= 0.5:
            scored.append((score, sheet_name))
    scored.sort(reverse=True)
    if scored and (len(scored) == 1 or scored[0][0] > scored[1][0]):
        return scored[0][1], workbook[scored[0][1]]
    raise ValueError(f"RR worksheet for {purpose} is missing or ambiguous; expected meaning similar to: {', '.join(names)}")


def find_header_row(ws, *required_terms, limit=100):
    wanted = [normalized_label(term) for term in required_terms]
    for row in range(1, min(ws.max_row, limit) + 1):
        labels = [normalized_label(ws.cell(row, column).value) for column in range(1, ws.max_column + 1)]
        if all(any(term in label for label in labels) for term in wanted):
            return row
    raise ValueError(f"Could not locate semantic header containing: {', '.join(required_terms)}")


def header_map(ws, row):
    result = {}
    for column in range(1, ws.max_column + 1):
        value = clean(ws.cell(row, column).value)
        if value is not None:
            result[re.sub(r"\s+", " ", str(value)).strip().lower()] = column
    return result


def matching_column(headers, *needles):
    for header, column in headers.items():
        if any(needle in header for needle in needles):
            return column
    return None


def source(sheet, row, column):
    return f"{sheet}!{get_column_letter(column)}{row}"


def row_fields(ws, sheet, row, columns):
    fields = {}
    for name, column in columns.items():
        if not column:
            continue
        value = clean(ws.cell(row, column).value)
        if value is not None:
            fields[name] = field(value, source(sheet, row, column))
    return fields


def populated_records(ws, sheet, header_row, start_row, columns, required=()):
    records = []
    for row in range(start_row, ws.max_row + 1):
        values = row_fields(ws, sheet, row, columns)
        if not values or any(name not in values for name in required):
            continue
        if any(str(item["value"]).strip().startswith("[") for item in values.values()):
            continue
        records.append({"sourceRow": row, "fields": values})
    return records


def choose_registration_layout(ws):
    # Current RRs label REG CODE in column A. Older accepted templates used B/C/D/E/F.
    for row in range(1, min(ws.max_row, 50) + 1):
        headers = header_map(ws, row)
        reg_code = matching_column(headers, "reg code", "registration type code")
        state = matching_column(headers, "activate", "status")
        admission_code = matching_column(headers, "admission item code")
        if reg_code and state and admission_code:
            return row, {
                "registration_code": reg_code,
                "registration_name": matching_column(headers, "reg type name", "registration type name"),
                "active": state,
                "admission_code": admission_code,
                "admission_name": matching_column(headers, "admission item") if admission_code != matching_column(headers, "admission item") else admission_code + 1,
                "admission_additional_text": matching_column(headers, "admission item additional text"),
                "group_registration": matching_column(headers, "register another person", "group registration"),
                "registration_path": matching_column(headers, "registration path"),
                "admission_description": matching_column(headers, "admission item description"),
                "badge_description": matching_column(headers, "badge description"),
                "registration_method": matching_column(headers, "registration method"),
                "reported": matching_column(headers, "reported"),
                "approval_required": matching_column(headers, "approval needed"),
                "pre_approval": matching_column(headers, "eligible for pre-approval"),
                "advanced_pre_registration": matching_column(headers, "advanced pre-registration"),
                "qualified": matching_column(headers, "qualified reg type"),
                "badge_color": matching_column(headers, "badge color"),
                "reprint_fee": matching_column(headers, "reprint fee"),
            }
    return 4, {
        "registration_code": 2, "registration_name": 3, "active": 4,
        "admission_code": 5, "admission_name": 6, "admission_description": 10,
    }


def tier_columns(ws, header_row, layout):
    excluded = {column for column in layout.values() if isinstance(column, int)}
    tiers = []
    for column in range(1, ws.max_column + 1):
        value = clean(ws.cell(header_row, column).value)
        if column in excluded or value is None:
            continue
        text = str(value)
        if re.search(r"(?:tier|insider|saver|early|advance|last chance|onsite)", text, re.I) or re.search(r"\d{1,2}/\d{1,2}", text):
            tiers.append((column, text))
    return tiers


def raw_rows(ws, sheet, start, end=None, columns=None):
    records = []
    for row in range(start, min(end or ws.max_row, ws.max_row) + 1):
        values = {}
        for column in (columns or range(1, ws.max_column + 1)):
            value = clean(ws.cell(row, column).value)
            if value is not None:
                values[get_column_letter(column)] = value
        if values:
            records.append({"sourceRow": row, "source": f"{sheet}!A{row}:{get_column_letter(ws.max_column)}{row}", "values": values})
    return records


def question_records(ws, sheet):
    header_row = find_header_row(ws, "question text", "question appearance", limit=80)
    headers = header_map(ws, header_row)
    columns = {
        "page": matching_column(headers, "page displayed"),
        "internal_name": matching_column(headers, "demo name", "internal name"),
        "respondent_scope": matching_column(headers, "company or individual"),
        "displayed_text": matching_column(headers, "question text"),
        "answer_code": matching_column(headers, "answer code"),
        "answer_text": matching_column(headers, "answer text", "answer selections"),
        "appearance": matching_column(headers, "question appearance"),
        "required": matching_column(headers, "required for registrant", "required"),
        "registration_type_visibility": matching_column(headers, "list reg types", "registration type visibility"),
        "determines_registration_type": matching_column(headers, "determine reg type"),
        "trigger_question": matching_column(headers, "trigger question"),
        "include_on_qr_code": matching_column(headers, "included on qr code", "include on qr code"),
        "notes": matching_column(headers, "notes"),
    }
    text_column = columns["displayed_text"]
    metadata_columns = [column for name, column in columns.items() if column and name not in {"displayed_text", "answer_code", "answer_text", "notes"}]
    definitions = []
    for row in range(header_row + 1, ws.max_row + 1):
        if text_column and clean(ws.cell(row, text_column).value) is not None and any(clean(ws.cell(row, column).value) is not None for column in metadata_columns):
            definitions.append(row)
    records = []
    for index, row in enumerate(definitions):
        next_row = definitions[index + 1] if index + 1 < len(definitions) else ws.max_row + 1
        fields = row_fields(ws, sheet, row, {name: column for name, column in columns.items() if name not in {"answer_code", "answer_text", "notes"}})
        choices = []
        continuation = []
        for answer_row in range(row + 1, next_row):
            code = clean(ws.cell(answer_row, columns["answer_code"]).value) if columns["answer_code"] else None
            text = clean(ws.cell(answer_row, columns["answer_text"]).value) if columns["answer_text"] else None
            if code is not None or text is not None:
                start = get_column_letter(columns["answer_code"] or columns["answer_text"])
                end = get_column_letter(columns["answer_text"] or columns["answer_code"])
                choices.append({"code": code, "text": text, "source": f"{sheet}!{start}{answer_row}:{end}{answer_row}"})
            extra = clean(ws.cell(answer_row, text_column).value) if text_column else None
            if extra is not None:
                continuation.append(str(extra))
        if choices:
            fields["ordered_answers"] = field(choices, f"{sheet}!{get_column_letter(columns['answer_code'])}{row + 1}:{get_column_letter(columns['answer_text'])}{next_row - 1}")
        if continuation and "displayed_text" in fields:
            fields["displayed_text"]["value"] = "\n\n".join([str(fields["displayed_text"]["value"]), *continuation])
            fields["displayed_text"]["source"] = f"{sheet}!{get_column_letter(text_column)}{row}:{get_column_letter(text_column)}{next_row - 1}"
        notes = clean(ws.cell(row, columns["notes"]).value) if columns["notes"] else None
        reference = clean(ws.cell(row, columns["internal_name"]).value) if columns["internal_name"] else None
        records.append({"sourceRow": row, "matchReference": reference, "fields": fields, "implementationNotes": notes})
    return records


wb = load_workbook(RR, data_only=True, read_only=False)
try:
    event_sheet, event_ws = worksheet_by_alias(wb, "event details", "Event Details", "Event Information", "Event Info", "Event Setup")
    links_sheet, links_ws = worksheet_by_alias(wb, "helpful and social links", "Helpful & Social Media Links", "Helpful Links", "Social Media Links")
    discount_sheet, discount_ws = worksheet_by_alias(wb, "discount codes", "Discount Code Template", "Discount Codes", "Discounts")
    question_sheet, question_ws = worksheet_by_alias(wb, "show questions", "Show Questions", "Registration Questions", "Questions")

    protected = {"event name", "event fp code", "event code"}
    event_fields = {}
    event_header = None
    for candidate in range(1, min(event_ws.max_row, 100) + 1):
        if any("general information" in normalized_label(event_ws.cell(candidate, column).value) for column in range(1, event_ws.max_column + 1)):
            event_header = candidate; break
    if event_header:
        headers = header_map(event_ws, event_header)
        label_column = matching_column(headers, "general information", "setting")
        notes_columns = [column for header, column in headers.items() if "note" in header or "additional info" in header]
        value_end = min(notes_columns) - 1 if notes_columns else event_ws.max_column
        event_rows = range(event_header + 1, min(event_ws.max_row, event_header + 140) + 1)
    else:
        label_column, value_end, event_rows = 1, event_ws.max_column, range(1, min(event_ws.max_row, 150) + 1)
    blank_run = 0
    for row in event_rows:
        key = clean(event_ws.cell(row, label_column).value)
        value_candidates = [(column, clean(event_ws.cell(row, column).value)) for column in range(label_column + 1, value_end + 1) if clean(event_ws.cell(row, column).value) is not None]
        if key is None and not value_candidates:
            blank_run += 1
            if event_header and blank_run >= 5: break
            continue
        blank_run = 0
        if key is None or not value_candidates:
            continue
        value_column, value = value_candidates[0]
        key_normalized = normalized_label(key)
        if key_normalized in protected or key_normalized.startswith("select one"):
            continue
        field_name = re.sub(r"[^a-z0-9]+", "_", str(key).lower()).strip("_")
        event_fields[field_name] = field(value, source(event_sheet, row, value_column))

    links = []
    footer_headers = []
    for row in range(1, min(links_ws.max_row, 100) + 1):
        headers = header_map(links_ws, row)
        if matching_column(headers, "footer options", "link label") and matching_column(headers, "visible") and matching_column(headers, "provide link"):
            footer_headers.append((row, headers))
    for header_index, (header_row, headers) in enumerate(footer_headers):
        end = footer_headers[header_index + 1][0] if header_index + 1 < len(footer_headers) else links_ws.max_row + 1
        label_column = matching_column(headers, "footer options", "link label")
        visible_column = matching_column(headers, "visible")
        target_column = matching_column(headers, "provide link", "url")
        context = " ".join(str(clean(links_ws.cell(candidate, 1).value) or "") for candidate in range(max(1, header_row - 3), header_row + 1)).lower()
        audience = "exhibitor" if "exhibitor" in context else "attendee_media"
        for row in range(header_row + 1, end):
            label = clean(links_ws.cell(row, label_column).value)
            if label and (normalized_label(label) == "countdown clock" or normalized_label(label).startswith("social media")):
                break
            visible = clean(links_ws.cell(row, visible_column).value)
            target = clean(links_ws.cell(row, target_column).value)
            if not label or normalized_label(label).startswith("for all") or (visible is None and target is None):
                continue
            links.append({"audience": audience, "label": label, "fields": {
                **({"visible": field(yes_no(visible), source(links_sheet, row, visible_column))} if visible is not None else {}),
                **({"target": field(target, source(links_sheet, row, target_column))} if target is not None else {}),
            }})
    countdown_messages = []
    countdown_found = False
    social = []
    for row in range(1, links_ws.max_row + 1):
        values = [(column, clean(links_ws.cell(row, column).value)) for column in range(1, links_ws.max_column + 1)]
        row_text = " ".join(str(value or "") for _, value in values).lower()
        if not countdown_found and normalized_label(clean(links_ws.cell(row, 1).value)) == "countdown clock" and row + 1 <= links_ws.max_row:
            countdown_found = True
            for item_row in range(row + 1, min(row + 12, links_ws.max_row + 1)):
                nonempty = [(column, value) for column, value in [(column, clean(links_ws.cell(item_row, column).value)) for column in range(1, links_ws.max_column + 1)] if value is not None]
                if len(nonempty) >= 2 and "social media" not in " ".join(str(value) for _, value in nonempty).lower():
                    countdown_messages.append({"name": nonempty[0][1], "dateRange": nonempty[1][1], "text": nonempty[2][1] if len(nonempty) > 2 else None, "source": f"{links_sheet}!{get_column_letter(nonempty[0][0])}{item_row}:{get_column_letter(nonempty[-1][0])}{item_row}"})
                elif "social media" in " ".join(str(value) for _, value in nonempty).lower():
                    break
        if "social media type" in row_text:
            headers = header_map(links_ws, row)
            network_column = matching_column(headers, "social media type", "network")
            visible_column = matching_column(headers, "visible")
            url_column = matching_column(headers, "url")
            for item_row in range(row + 1, links_ws.max_row + 1):
                network = clean(links_ws.cell(item_row, network_column).value)
                visible = clean(links_ws.cell(item_row, visible_column).value) if visible_column else None
                url = clean(links_ws.cell(item_row, url_column).value) if url_column else None
                if network and (visible is not None or url is not None):
                    social.append({"network": network, "visible": yes_no(visible), "url": url, "source": f"{links_sheet}!{get_column_letter(network_column)}{item_row}:{get_column_letter(url_column)}{item_row}"})
            break

    reg_sheet, reg_ws = worksheet_by_alias(wb, "registration types and pricing", "NEW Reg Types & Pricing", "Reg Types & Pricing")
    reg_header, layout = choose_registration_layout(reg_ws)
    tiers = tier_columns(reg_ws, reg_header, layout)
    active_rows = []
    for row in range(reg_header + 1, reg_ws.max_row + 1):
        status = clean(reg_ws.cell(row, layout["active"]).value)
        code = clean(reg_ws.cell(row, layout["registration_code"]).value)
        name = clean(reg_ws.cell(row, layout.get("registration_name") or 0).value) if layout.get("registration_name") else None
        if str(status or "").upper() not in {"ACTIVATE", "REQUIRED", "ACTIVE", "YES"} or not (code or name):
            continue
        fields = row_fields(reg_ws, reg_sheet, row, layout)
        active_rows.append({"sourceRow": row, "fields": fields})

    reg_types_by_key, admissions_by_key, prices = {}, {}, []
    paths = {}
    for item in active_rows:
        row, fields = item["sourceRow"], item["fields"]
        reg_code = fields.get("registration_code", {}).get("value")
        reg_name = fields.get("registration_name", {}).get("value")
        reg_key = str(reg_code or reg_name)
        reg_fields = {key: value for key, value in fields.items() if key not in {"admission_code", "admission_name", "admission_description", "admission_additional_text", "badge_description"}}
        existing = reg_types_by_key.setdefault(reg_key, {"matchReference": reg_code or reg_name, "sourceRows": [], "fields": reg_fields})
        existing["sourceRows"].append(row)
        path_ref = fields.get("registration_path", {}).get("value")
        if path_ref:
            path = paths.setdefault(str(path_ref), {"matchReference": path_ref, "registrationTypes": [], "sourceEvidence": []})
            path["registrationTypes"].append(reg_code or reg_name)
            evidence_columns = [layout["registration_path"], layout.get("registration_code") or layout.get("registration_name")]
            evidence_range = f"{reg_sheet}!{get_column_letter(min(evidence_columns))}{row}:{get_column_letter(max(evidence_columns))}{row}"
            path["sourceEvidence"].append(field({"path": path_ref, "registrationType": reg_code or reg_name}, evidence_range))
        admission_code = fields.get("admission_code", {}).get("value")
        admission_name = fields.get("admission_name", {}).get("value")
        if admission_code or admission_name:
            admission_key = str(admission_code or admission_name)
            admission = admissions_by_key.setdefault(admission_key, {
                "matchReference": admission_code or admission_name, "sourceRows": [],
                "fields": {key: value for key, value in fields.items() if key in {"admission_code", "admission_name", "admission_description", "admission_additional_text", "badge_description"}},
                "registrationTypes": [], "registrationTypeEvidence": [],
            })
            admission["sourceRows"].append(row)
            admission["registrationTypes"].append(reg_code or reg_name)
            association_columns = [column for column in (layout.get("registration_code") or layout.get("registration_name"), layout.get("admission_code") or layout.get("admission_name")) if column]
            association_source = f"{reg_sheet}!{get_column_letter(min(association_columns))}{row}:{get_column_letter(max(association_columns))}{row}"
            admission["registrationTypeEvidence"].append(field({"registrationType": reg_code or reg_name, "admissionItem": admission_code or admission_name}, association_source))
        tier_values = []
        for column, label in tiers:
            value = clean(reg_ws.cell(row, column).value)
            if value is not None:
                tier_values.append({"nameOrDateRange": label, "value": value, "source": source(reg_sheet, row, column)})
        if tier_values:
            prices.append({"registrationType": reg_code or reg_name, "admissionItem": admission_code or admission_name, "tiers": tier_values, "sourceRow": row})

    discount_header = find_header_row(discount_ws, "discount code", "method")
    discount_headers = header_map(discount_ws, discount_header)
    discount_columns = {
        "name": matching_column(discount_headers, "name"), "code": matching_column(discount_headers, "discount code"),
        "type": matching_column(discount_headers, "discount type"), "method": matching_column(discount_headers, "method"),
        "amount_or_percentage": matching_column(discount_headers, "amount/percentage"),
        "effective_from": matching_column(discount_headers, "effective from"), "effective_to": matching_column(discount_headers, "effective to"),
        "capacity": matching_column(discount_headers, "capacity"), "stackable": matching_column(discount_headers, "stackable"),
        "usable_by": matching_column(discount_headers, "can be used by"), "count_guests": matching_column(discount_headers, "counts guests"),
        "active": matching_column(discount_headers, "active"), "internal_note": matching_column(discount_headers, "internal note"),
        "admission_items": matching_column(discount_headers, "admission items"), "tracks": matching_column(discount_headers, "tracks"),
        "sessions": matching_column(discount_headers, "sessions"),
    }
    discounts = populated_records(discount_ws, discount_sheet, discount_header, discount_header + 1, discount_columns, required=("name", "code"))
    discounts = [item for item in discounts if str(item["fields"].get("active", {}).get("value", "")).strip().lower() == "yes"]
    discount_import = write_discount_import(JOB_DIR / "discount-import.xlsx", discount_ws, discount_columns, discounts)

    optional_items = []
    if "Optional Items" in wb.sheetnames:
        optional_ws = wb["Optional Items"]
        optional_headers = header_map(optional_ws, 2)
        optional_columns = {re.sub(r"[^a-z0-9]+", "_", name).strip("_"): column for name, column in optional_headers.items()}
        optional_items = populated_records(optional_ws, "Optional Items", 2, 3, optional_columns, required=("item_code", "item_title"))

    sessions = []
    if "Sessions" in wb.sheetnames:
        session_ws = wb["Sessions"]
        session_headers = header_map(session_ws, 1)
        session_columns = {re.sub(r"[^a-z0-9]+", "_", name).strip("_"): column for name, column in session_headers.items()}
        sessions = populated_records(session_ws, "Sessions", 1, 2, session_columns, required=("item_code_sess_xxx_in_sessionboard_or_whatever_code_you_want_to_use_if_not_using_sessionboard",))

    group_discounts = []
    if "Group_Volume Discounts" in wb.sheetnames:
        group_ws = wb["Group_Volume Discounts"]
        group_headers = header_map(group_ws, 2)
        group_columns = {re.sub(r"[^a-z0-9]+", "_", name).strip("_"): column for name, column in group_headers.items()}
        group_discounts = populated_records(group_ws, "Group_Volume Discounts", 2, 3, group_columns, required=("name",))

    questions = question_records(question_ws, question_sheet)
    inline_content_links = []
    seen_inline_links = set()
    for question in questions:
        for field_name, evidence in question.get("fields", {}).items():
            value = evidence.get("value") if isinstance(evidence, dict) else None
            if not isinstance(value, str):
                continue
            for match in re.findall(r'https?://[^\s<>"\']+', value, re.IGNORECASE):
                url = match.rstrip(".,;:)]}")
                key = (evidence.get("source"), url)
                if key in seen_inline_links:
                    continue
                seen_inline_links.add(key)
                inline_content_links.append({
                    "label": f"Inline link in {question.get('matchReference') or field_name}",
                    "target": field(url, evidence["source"]),
                })
    policies = []
    if "Policies & Rules" in wb.sheetnames:
        policy_ws = wb["Policies & Rules"]
        for row in range(3, policy_ws.max_row + 1):
            setting, value, notes = (clean(policy_ws.cell(row, col).value) for col in (1, 2, 3))
            if setting and value is not None:
                policies.append({"setting": setting, "value": value, "notes": notes, "source": f"Policies & Rules!A{row}:C{row}"})

    integration_requirements = raw_rows(wb["Integrations"], "Integrations", 3, 18, range(1, 5)) if "Integrations" in wb.sheetnames else []
    communication_requirements = raw_rows(wb["Communications"], "Communications", 3, 20, range(1, 5)) if "Communications" in wb.sheetnames else []
    badge_requirements = raw_rows(wb["Badge & Ticket Layouts"], "Badge & Ticket Layouts", 3, 33) if "Badge & Ticket Layouts" in wb.sheetnames else []
    scan_go_requirements = raw_rows(wb["Onsite Scan & Go Screen"], "Onsite Scan & Go Screen", 3, 11, range(5, 8)) if "Onsite Scan & Go Screen" in wb.sheetnames else []
    approval_requirements = raw_rows(wb["Approvals"], "Approvals", 4, 24, range(1, 4)) if "Approvals" in wb.sheetnames else []
    capability_gaps = []
    if approval_requirements and not any(record["values"].get("A") and record["values"].get("B") for record in approval_requirements):
        capability_gaps.append({
            "domain": "associations", "missingCapability": "none",
            "reason": "The Approvals worksheet contains guidance text but no business-type association to configure.",
        })

    domains = {
        "event_settings": {"fields": event_fields, "policies": policies},
        "site_designer": {"footerLinks": links, "countdownMessages": countdown_messages, "socialLinks": social, "inlineContentLinks": inline_content_links},
        "registration_paths": {"items": list(paths.values())},
        "registration_types": {"items": list(reg_types_by_key.values())},
        "admission_items": {"items": list(admissions_by_key.values())},
        "optional_items": {"items": optional_items},
        "pricing": {"tierHeaders": [{"nameOrDateRange": label, "source": source(reg_sheet, reg_header, column)} for column, label in tiers], "items": prices},
        "discounts_vouchers": {"discounts": discounts, "groupDiscounts": group_discounts},
        "questions": {"items": questions},
        "sessions": {"items": sessions},
        "integrations": {"requirements": integration_requirements},
        "communications": {"requirements": communication_requirements, "constraint": "Event-scoped templates and triggered-confirmation settings may be configured; never manually send, test-send, schedule a blast, or access recipients."},
        "badges_onsite": {"badgeRequirements": badge_requirements, "scanAndGoText": scan_go_requirements},
        "associations": {"approvalRequirements": approval_requirements},
        "final_qa": {"checks": ["selected event identity unchanged", "event remains unpublished", "all RR-requested event configuration reread", "no unintended duplicates", "no communications sent", "no attendee/contact access"]},
    }

    def count_fields(item):
        if isinstance(item, dict):
            if "value" in item and "source" in item:
                return 1
            return sum(count_fields(value) for value in item.values())
        if isinstance(item, list):
            return sum(count_fields(value) for value in item)
        return 0

    result = {
        "schemaVersion": 4,
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "rr": {"path": str(RR), "sha256": hashlib.sha256(RR.read_bytes()).hexdigest(), "authority": "uploaded_rr"},
        "target": TARGET,
        "domains": domains,
        "protected": [
            "selected event identity", "publish/go live", "delete/archive", "communications",
            "attendees/contacts", "account-global or reusable definitions",
        ],
        "artifacts": {"discountImport": discount_import},
        "capabilityGaps": capability_gaps,
        "counts": {
            "applicableFields": count_fields(domains),
            "registrationTypeRecords": len(reg_types_by_key), "admissionRecords": len(admissions_by_key),
            "pricingRecords": len(prices), "discountRecords": len(discounts),
            "groupDiscountRecords": len(group_discounts), "questionRecords": len(questions),
            "optionalItemRecords": len(optional_items), "sessionRecords": len(sessions),
            "siteLinkRecords": len(links) + len(social) + len(inline_content_links), "integrationRequirementRows": len(integration_requirements),
            "communicationRequirementRows": len(communication_requirements), "badgeOnsiteRequirementRows": len(badge_requirements) + len(scan_go_requirements),
        },
    }
    atomic_json(OUT, result)
    print(json.dumps({"ok": True, "output": str(OUT), "counts": result["counts"], "capabilityGaps": capability_gaps}, indent=2))
finally:
    wb.close()
