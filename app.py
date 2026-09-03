from __future__ import annotations
import asyncio, json, os, re, shutil, signal, subprocess, threading, time
from datetime import date, datetime, time as datetime_time, timezone
from pathlib import Path
from urllib.parse import urlparse
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse
from browser_gate import initialize as initialize_gate, read as read_gate, write as write_gate, request_user, shield_agent, lock_file
from browser_runtime import RUNTIME as BROWSER_RUNTIME_PATH, initialize as initialize_browser_runtime, probe as probe_browser_runtime, local_probe as local_runtime_probe, load as load_browser_runtime, tool_probe
from scope_manifest import SCOPE_MANIFEST, SCOPE_WORKBOOK, load_manifest as load_scope_manifest

ROOT=Path(__file__).resolve().parent
DATA=ROOT/'data'; CURRENT=DATA/'current'; RUNS=DATA/'runs'
STATE=CURRENT/'state.json'; LOG=CURRENT/'activity.log'; REPORT=CURRENT/'final-report.json'
AUTH_SETTINGS=DATA/'auth-settings.json'
AUTHORIZED_EVENT_NAME='(C+D) Medtrade Testing Clone 2'
RUN_MODE='mock'
app=FastAPI(title='CVENT Agent')
_proc: subprocess.Popen|None=None
_lock=threading.Lock()
_workbook_lock=threading.Lock()

def now(): return datetime.now(timezone.utc).isoformat()
def read_json(path,default):
    try: return json.loads(path.read_text())
    except Exception: return default
def atomic_json(path,data):
    path.parent.mkdir(parents=True,exist_ok=True); tmp=path.with_suffix('.tmp'); tmp.write_text(json.dumps(data,indent=2)); tmp.replace(path)
def append_log(msg):
    CURRENT.mkdir(parents=True,exist_ok=True)
    with LOG.open('a') as f: f.write(f'{now()}  {msg}\n')
def product_facing(value):
    if value=='PI_EGO': return 'CVENT_EGO'
    if isinstance(value,str): return re.sub(r'\bpi(?:\s+agent)?\b','CVENT Agent',value.replace(str(ROOT),'[CVENT Agent workspace]'),flags=re.I)
    if isinstance(value,list): return [product_facing(item) for item in value]
    if isinstance(value,dict): return {key:product_facing(item) for key,item in value.items()}
    return value
def fresh_state(filename=None):
    t=now(); return {'status':'ready' if filename else 'waiting_for_rr','current_stage':'upload','current_action':f'Ready — mock RR can modify only {AUTHORIZED_EVENT_NAME} within Intake Emerald scope' if filename else 'Upload mock RR workbook','completed':[],'pending':['target_discovery','event_basics','theme_branding','header_footer_body','registration_paths','registration_types','admission_items','pricing_fees','discounts','registration_questions','terms_policies','final_qa'],'review_required':[],'rr_file':filename,'run_mode':RUN_MODE,'authorized_event_name':AUTHORIZED_EVENT_NAME,'target_url':'','target_identity':AUTHORIZED_EVENT_NAME,'started_at':None,'process_started_at':None,'last_run_seconds':0,'updated_at':t,'pi_pid':None,'pi_session':None}
def auth_settings():
    saved=read_json(AUTH_SETTINGS,{})
    cookie_store=DATA/'steel-profile-local'/'Default'/'Cookies'
    return {'organization_id':saved.get('organization_id',''),'authenticated_at':saved.get('authenticated_at'),'cookies_saved':bool(saved.get('authenticated_at') and cookie_store.exists()),'microsoft_sso_persistent':bool(saved.get('microsoft_sso_persistent')),'cvent_cookie_count':saved.get('cvent_cookie_count',0),'microsoft_cookie_count':saved.get('microsoft_cookie_count',0),'cookie_store':str(cookie_store)}
def ensure():
    CURRENT.mkdir(parents=True,exist_ok=True); RUNS.mkdir(parents=True,exist_ok=True); (DATA/'chrome-profile').mkdir(parents=True,exist_ok=True)
    if not STATE.exists(): atomic_json(STATE,fresh_state('input.xlsx' if (CURRENT/'input.xlsx').exists() else None))
    LOG.touch(exist_ok=True)
    if not (CURRENT/'browser-gate.json').exists(): initialize_gate()
    if (CURRENT/'input.xlsx').exists() and not REPORT.exists(): atomic_json(REPORT,{'status':'INCOMPLETE','unresolved_items':['Build not started'],'real_reads':[],'real_writes':[],'guardrails':{'published':0,'emails_sent':0,'deletes':0,'global_mutations':0},'updated_at':now()})
def steel_command(command,url=None,timeout=75):
    cmd=[str(ROOT/'steel_session.py'),command]
    if url: cmd += ['--url',url]
    result=subprocess.run(cmd,cwd=ROOT,text=True,capture_output=True,timeout=timeout)
    for line in reversed(result.stdout.splitlines()):
        if line.startswith('STEEL_RESULT='):
            data=json.loads(line.split('=',1)[1]); break
    else: data={'provider':'steel','running':False,'error':(result.stderr or result.stdout or 'Steel command failed')[-1000:]}
    if result.returncode and not data.get('error'): data['error']=f'Steel command exited {result.returncode}'
    return data
def chrome_status():
    data=steel_command('status',timeout=20)
    return {k:data.get(k) for k in ('provider','running','status','id','profile_id','viewer_url','error','login_required')}
def chrome_auth(): return steel_command('page',timeout=60)
def chrome_login_required(): return chrome_auth().get('auth_status')!='authenticated'
def launch_chrome(url='https://app.cvent.com/'):
    data=steel_command('ensure',timeout=90)
    if not data.get('running'): return data
    page=steel_command('page',url=url,timeout=75)
    data['login_required']=page.get('login_required'); data['url']=page.get('url'); data['title']=page.get('title')
    return data
def ensure_browser_runtime(full_probe=False):
    try:
        runtime=load_browser_runtime(BROWSER_RUNTIME_PATH); local_runtime_probe(runtime)
    except Exception:
        initialize_gate(); runtime=initialize_browser_runtime()
    if full_probe:probe_browser_runtime(BROWSER_RUNTIME_PATH,full=True)
    return load_browser_runtime(BROWSER_RUNTIME_PATH)
def valid_target(url):
    try:
        u=urlparse(url); host=(u.hostname or '').lower()
        return u.scheme=='https' and (host=='cvent.com' or host.endswith('.cvent.com')) and len(u.path+u.query)>1
    except Exception: return False
def authorized_target_url():
    lock=read_json(CURRENT/'authorized-target.json',{})
    url=lock.get('url','');key=event_key_from_url(url)
    return url if lock.get('name')==AUTHORIZED_EVENT_NAME and valid_target(url) and key and lock.get('event_key')==key else ''
def archive_current():
    if not CURRENT.exists() or not any(CURRENT.iterdir()): return
    stamp=datetime.now().strftime('%Y%m%d-%H%M%S')
    dest=RUNS/stamp; n=1
    while dest.exists(): n+=1; dest=RUNS/f'{stamp}-{n}'
    shutil.move(str(CURRENT),str(dest))

def job_pid():
    if _proc is not None and _proc.poll() is None:return _proc.pid
    pid=read_json(STATE,{}).get('pi_pid')
    if not isinstance(pid,int):return None
    try:
        command=subprocess.check_output(['ps','-p',str(pid),'-o','command='],text=True).strip()
        return pid if command and Path(command.split()[0]).name=='pi' else None
    except Exception:return None
def running(): return job_pid() is not None
def process_tree(root):
    rows=[]
    try:
        for line in subprocess.check_output(['ps','-axo','pid=,ppid='],text=True).splitlines():
            pid,ppid=map(int,line.split()); rows.append((pid,ppid))
    except Exception: rows=[]
    children=[]
    def walk(parent):
        for pid,ppid in rows:
            if ppid==parent: walk(pid); children.append(pid)
    walk(root); return children+[root]
def stop_process_tree(root):
    pids=process_tree(root)
    for sig in (signal.SIGTERM,signal.SIGKILL):
        for pid in pids:
            try: os.kill(pid,sig)
            except ProcessLookupError: pass
            except PermissionError: pass
        time.sleep(.7)

def render_prompt():
    vals={'RR_PATH':str((CURRENT/'input.xlsx').resolve()),'SCOPE_WORKBOOK_PATH':str(SCOPE_WORKBOOK.resolve()),'SCOPE_MANIFEST_PATH':str(SCOPE_MANIFEST.resolve()),'TARGET_URL':f'DISCOVER EXACTLY {AUTHORIZED_EVENT_NAME} — THE RR IS MOCK INPUT AND MUST NOT SELECT THE TARGET','STATE_PATH':str(STATE.resolve()),'LOG_PATH':str(LOG.resolve()),'REPORT_PATH':str(REPORT.resolve()),'AUTH_SETTINGS_PATH':str(AUTH_SETTINGS.resolve()),'BROWSER_RUNTIME_PATH':str(BROWSER_RUNTIME_PATH.resolve()),'BROWSER_TOOL_PATH':str((ROOT/'browser_tool.py').resolve()),'STATUS_HELPER':str((ROOT/'status_update.py').resolve())}
    text=(ROOT/'PI_PROMPT.md').read_text()
    for k,v in vals.items(): text=text.replace('{{'+k+'}}',v)
    (CURRENT/'job-prompt.md').write_text(text); return text

def spawn_pi(message,target=None,resume=False):
    global _proc
    sessions=CURRENT/'pi-sessions'; sessions.mkdir(parents=True,exist_ok=True)
    cmd=['pi','-p','--approve','--no-extensions','--no-skills','--skill',str(ROOT/'.agents/skills/cvent-browser/SKILL.md'),'--no-prompt-templates','--no-context-files','--session-dir',str(sessions),'--name','cvent-one-shot']
    if resume:
        files=sorted(sessions.glob('*.jsonl'),key=lambda p:p.stat().st_mtime,reverse=True)
        if not files: raise HTTPException(409,'No CVENT Agent session to continue')
        cmd += ['--session',str(files[0]),message]
    else: cmd += [message]
    out=open(CURRENT/'pi-output.log','a',buffering=1)
    _proc=subprocess.Popen(cmd,cwd=ROOT,stdout=out,stderr=subprocess.STDOUT,text=True,start_new_session=True)
    st=read_json(STATE,fresh_state('input.xlsx')); st.update({'status':'running','current_stage':'starting','current_action':'CVENT Agent is starting','target_url':target or st.get('target_url',''),'pi_pid':_proc.pid,'started_at':st.get('started_at') or now(),'process_started_at':now(),'updated_at':now()}); atomic_json(STATE,st)
    append_log(('Resuming' if resume else 'Started')+f' CVENT Agent process PID {_proc.pid}')
    threading.Thread(target=monitor_pi,args=(_proc,out),daemon=True).start()
    return _proc.pid

def monitor_pi(proc,out):
    global _proc
    code=proc.wait(); out.close(); time.sleep(.2)
    st=read_json(STATE,{})
    sessions=sorted((CURRENT/'pi-sessions').glob('*.jsonl'),key=lambda p:p.stat().st_mtime,reverse=True) if (CURRENT/'pi-sessions').exists() else []
    if sessions: st['pi_session']=str(sessions[0])
    if st.get('status')=='running':
        st['status']='agent_stopped' if code==0 else 'failed'; st['current_action']='CVENT Agent stopped before a final verdict' if code==0 else f'CVENT Agent exited with code {code}'
        append_log(st['current_action'])
    if st.get('process_started_at'):
        try:st['last_run_seconds']=max(0,int((datetime.now(timezone.utc)-datetime.fromisoformat(st['process_started_at'])).total_seconds()))
        except Exception:pass
    st['process_started_at']=None;st['pi_pid']=None;st['updated_at']=now(); atomic_json(STATE,st)
    with _lock:
        if _proc is proc: _proc=None

@app.on_event('startup')
def startup():
    # Keep the control UI available while Steel and Pi are stopped.
    ensure()

@app.get('/',response_class=HTMLResponse)
def home(): return HTMLResponse((ROOT/'templates/index.html').read_text(),headers={'Cache-Control':'no-store, no-cache, must-revalidate','Pragma':'no-cache','Expires':'0'})

@app.get('/steel-viewer',response_class=HTMLResponse)
def steel_viewer():
    try:
        import urllib.request
        with urllib.request.urlopen('http://127.0.0.1:3005/v1/sessions/debug',timeout=10) as r: html=r.read().decode('utf-8')
        html=html.replace('ws://0.0.0.0:3000','ws://127.0.0.1:3005').replace('http://0.0.0.0:3000','http://127.0.0.1:3005')
        safety="""<script>(()=>{let user=false;const stop=e=>{if(!user){e.preventDefault();e.stopImmediatePropagation();try{document.activeElement?.blur()}catch{}}};['pointerdown','pointerup','pointermove','mousedown','mouseup','mousemove','click','dblclick','contextmenu','wheel','touchstart','touchmove','touchend','keydown','keyup','keypress','focusin'].forEach(n=>document.addEventListener(n,stop,{capture:true,passive:false}));async function sync(){try{const r=await fetch('/api/browser/ownership',{cache:'no-store'}),d=await r.json();user=d.ownership==='USER'&&d.desiredOwnership==='USER';document.documentElement.dataset.controlOwner=user?'USER':'AGENT';document.body.style.pointerEvents=user?'auto':'none';if(!user)try{document.activeElement?.blur()}catch{}}catch{user=false;document.body.style.pointerEvents='none'}}sync();setInterval(sync,400)})()</script>"""
        html=html.replace('</body>',safety+'</body>')
        return HTMLResponse(html,headers={'Cache-Control':'no-store'})
    except Exception as e: raise HTTPException(503,f'Local Steel viewer unavailable: {e}')

@app.get('/api/browser/ownership')
def browser_ownership():
    return JSONResponse(read_gate(),headers={'Cache-Control':'no-store'})

@app.get('/api/status')
def status():
    ensure(); st=read_json(STATE,fresh_state()); st['run_mode']=RUN_MODE; st['authorized_event_name']=AUTHORIZED_EVENT_NAME; st['browser_strategy']='EGO DIRECT · SAME STEEL RUNTIME'; st['browser_gate']=read_gate();
    try:
        runtime=load_browser_runtime(BROWSER_RUNTIME_PATH); st['browser_runtime']={k:runtime.get(k) for k in ('browserRuntimeId','apiOrigin','cdpEndpoint','viewerUrl','targetBrowserIdentity','verifiedAt')}
    except Exception: st['browser_runtime']=None
    try:
        scope=load_scope_manifest();st['automation_scope']={'valid':True,'authority':scope['authority'],'counts':scope['counts'],'sha256':scope['sourceSha256']}
    except Exception as e:st['automation_scope']={'valid':False,'error':str(e)}
    st['activity_log']=LOG.read_text(errors='replace').splitlines()[-200:]; st['final_report']=read_json(REPORT,None); st['browser']=chrome_status(); st['browser'].pop('id',None); st['auth_settings']=auth_settings(); st['auth_settings'].pop('cookie_store',None); st['agent_process_running']=running(); st['agent_pid']=st.get('pi_pid'); st['agent_session_saved']=bool(st.get('pi_session'))
    path=CURRENT/'input.xlsx';st['rr_version']=str(path.stat().st_mtime_ns) if path.exists() else None
    if st['agent_process_running'] and st.get('process_started_at'):
        try: st['elapsed_seconds']=max(0,int((datetime.now(timezone.utc)-datetime.fromisoformat(st['process_started_at'])).total_seconds()))
        except Exception: st['elapsed_seconds']=0
    else:
        st['elapsed_seconds']=0;st['agent_pid']=None
    st.pop('pi_pid',None);st.pop('pi_session',None)
    return JSONResponse(product_facing(st),headers={'Cache-Control':'no-store'})

@app.get('/api/scope')
def automation_scope():
    try:return JSONResponse(load_scope_manifest(),headers={'Cache-Control':'no-store'})
    except Exception as e:raise HTTPException(500,f'Automation scope is invalid: {e}')

@app.get('/api/workbook')
def workbook_info():
    path=CURRENT/'input.xlsx'
    if not path.exists(): raise HTTPException(404,'No RR workbook uploaded')
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=False)
    try: sheets=[{'name':ws.title,'rows':ws.max_row,'columns':ws.max_column} for ws in wb.worksheets]
    finally: wb.close()
    return JSONResponse({'file':read_json(STATE,{}).get('rr_file') or path.name,'version':str(path.stat().st_mtime_ns),'sheets':sheets},headers={'Cache-Control':'no-store'})

@app.get('/api/workbook/sheet')
def workbook_sheet(name:str,start:int=1,limit:int=80):
    path=CURRENT/'input.xlsx'
    if not path.exists(): raise HTTPException(404,'No RR workbook uploaded')
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils import get_column_letter
    wb=load_workbook(path,read_only=False,data_only=False)
    try:
        if name not in wb.sheetnames: raise HTTPException(404,'Worksheet not found')
        ws=wb[name]; start=max(1,start); limit=max(10,min(limit,150)); end=min(ws.max_row,start+limit-1); width=min(ws.max_column,60)
        def value(v):
            if v is None:return ''
            if hasattr(v,'isoformat'):return v.isoformat()
            return str(v)
        rows=[[value(ws.cell(r,c).value) for c in range(1,width+1)] for r in range(start,end+1)]
        editable=[[not ws.protection.sheet and not isinstance(ws.cell(r,c),MergedCell) for c in range(1,width+1)] for r in range(start,end+1)]
        return JSONResponse({'name':name,'version':str(path.stat().st_mtime_ns),'start':start,'end':end,'total_rows':ws.max_row,'total_columns':ws.max_column,'columns':[get_column_letter(c) for c in range(1,width+1)],'rows':rows,'editable':editable,'protected':bool(ws.protection.sheet)},headers={'Cache-Control':'no-store'})
    finally: wb.close()

def edited_cell_value(cell,text):
    if text=='':return None
    if text.startswith('='):return text
    current=cell.value
    try:
        if isinstance(current,bool):
            lowered=text.strip().lower()
            if lowered not in ('true','false'):raise ValueError
            return lowered=='true'
        if isinstance(current,int) and not isinstance(current,bool):return int(text)
        if isinstance(current,float):return float(text)
        if isinstance(current,datetime):return datetime.fromisoformat(text)
        if isinstance(current,date):return date.fromisoformat(text)
        if isinstance(current,datetime_time):return datetime_time.fromisoformat(text)
    except ValueError:raise HTTPException(400,f'Value {text!r} is invalid for {cell.coordinate}')
    return text

def reset_after_workbook_edit(filename):
    for name in ('benchmark-results.json','build-checklist.json','domain-results.json','expected-domains.json','input.inspection.json','input.inspection-summary.json','review-required.json','rr-checklist.json','rr-execution-checklist.json','rr-focus.json','rr-focused.txt','rr-inspection.json','rr-inspection-summary.json','rr-question-checklist.json'):
        try:(CURRENT/name).unlink()
        except FileNotFoundError:pass
    st=fresh_state(filename);locked=authorized_target_url()
    if locked:st.update({'target_url':locked,'target_identity':AUTHORIZED_EVENT_NAME})
    st['current_action']='RR workbook edited — ready to reread requirements';atomic_json(STATE,st)
    atomic_json(REPORT,{'status':'INCOMPLETE','unresolved_items':['Workbook changed; build must reread requirements'],'real_reads':[],'real_writes':[],'guardrails':{'published':0,'emails_sent':0,'deletes':0,'global_mutations':0},'updated_at':now()})

@app.patch('/api/workbook')
def update_workbook(payload:dict):
    path=CURRENT/'input.xlsx'
    if not path.exists():raise HTTPException(404,'No RR workbook uploaded')
    if running():raise HTTPException(409,'Stop CVENT Agent before editing the RR workbook')
    changes=payload.get('changes');version=payload.get('version')
    if not isinstance(changes,list) or not changes or len(changes)>2000:raise HTTPException(400,'Submit between 1 and 2000 cell changes')
    with _workbook_lock:
        if running():raise HTTPException(409,'Stop CVENT Agent before editing the RR workbook')
        if str(version)!=str(path.stat().st_mtime_ns):raise HTTPException(409,'The workbook changed; reload it before saving')
        from openpyxl import load_workbook
        from openpyxl.cell.cell import MergedCell
        wb=load_workbook(path,data_only=False)
        try:
            normalized=[];seen=set()
            for change in changes:
                if not isinstance(change,dict):raise HTTPException(400,'Each cell change must be an object')
                sheet=change.get('sheet');row=change.get('row');column=change.get('column');value=change.get('value')
                if not isinstance(sheet,str) or sheet not in wb.sheetnames:raise HTTPException(400,'Unknown worksheet')
                if not isinstance(row,int) or not 1<=row<=1048576 or not isinstance(column,int) or not 1<=column<=16384:raise HTTPException(400,'Invalid cell coordinates')
                if not isinstance(value,str) or len(value)>32767:raise HTTPException(400,'Cell values must be text no longer than 32,767 characters')
                key=(sheet,row,column)
                if key in seen:raise HTTPException(400,'Duplicate cell change')
                seen.add(key);ws=wb[sheet];cell=ws.cell(row,column)
                if ws.protection.sheet:raise HTTPException(409,f'Worksheet {sheet} is protected')
                if isinstance(cell,MergedCell):raise HTTPException(409,f'{sheet}!{cell.coordinate} is a non-editable merged cell')
                normalized.append((cell,edited_cell_value(cell,value)))
            backups=CURRENT/'workbook-backups';backups.mkdir(parents=True,exist_ok=True)
            stamp=datetime.now().strftime('%Y%m%d-%H%M%S-%f');backup=backups/f'input-{stamp}.xlsx';shutil.copy2(path,backup)
            for cell,value in normalized:cell.value=value
            temp=path.with_name('input.editing.xlsx');wb.save(temp);os.replace(temp,path)
        finally:
            wb.close()
            try:temp.unlink()
            except (NameError,FileNotFoundError):pass
        filename=read_json(STATE,{}).get('rr_file') or 'input.xlsx';reset_after_workbook_edit(filename);append_log(f'Saved {len(normalized)} RR workbook cell edit(s); requirements must be reread before build')
        return {'ok':True,'saved':len(normalized),'version':str(path.stat().st_mtime_ns),'backup':backup.name}

@app.post('/api/upload')
def upload(rr:UploadFile=File(...)):
    ensure()
    if running(): raise HTTPException(409,'CVENT Agent is running')
    name=rr.filename or ''
    if not name.lower().endswith('.xlsx'): raise HTTPException(400,'Upload an .xlsx file')
    archive_current(); CURRENT.mkdir(parents=True,exist_ok=True)
    with (CURRENT/'input.xlsx').open('wb') as f: shutil.copyfileobj(rr.file,f)
    atomic_json(STATE,fresh_state(name)); LOG.write_text('')
    atomic_json(REPORT,{'status':'INCOMPLETE','unresolved_items':['Build not started'],'real_reads':[],'real_writes':[],'guardrails':{'published':0,'emails_sent':0,'deletes':0,'global_mutations':0},'updated_at':now()})
    append_log(f'Uploaded RR workbook: {name}')
    return {'ok':True,'file':name}

async def capture_login_facts():
    import urllib.request, websockets
    pages=json.load(urllib.request.urlopen('http://127.0.0.1:9334/json/list',timeout=5))
    page=next((x for x in pages if x.get('type')=='page' and 'cvent.com' in (x.get('url') or '')),None)
    if not page: raise RuntimeError('No Cvent page is open in Steel')
    ws=page['webSocketDebuggerUrl'].replace('ws://127.0.0.1/','ws://127.0.0.1:9334/').replace('ws://localhost/','ws://127.0.0.1:9334/')
    async with websockets.connect(ws,origin='http://127.0.0.1:9334',open_timeout=10) as socket:
        await socket.send(json.dumps({'id':1,'method':'Network.getAllCookies'}))
        while True:
            reply=json.loads(await asyncio.wait_for(socket.recv(),10))
            if reply.get('id')==1: break
    cookies=reply.get('result',{}).get('cookies',[])
    org=next((c.get('value','') for c in cookies if c.get('name')=='org-id' and c.get('domain','').endswith('cvent.com')), '')
    return {'organization_id':org,'microsoft_sso_persistent':any(c.get('name')=='ESTSAUTHPERSISTENT' for c in cookies),'cvent_cookie_count':sum(c.get('domain','').endswith('cvent.com') for c in cookies),'microsoft_cookie_count':sum('microsoftonline.com' in c.get('domain','') for c in cookies)}

@app.post('/api/auth-settings')
def save_auth_settings():
    ensure(); auth=chrome_auth()
    if auth.get('auth_status')!='authenticated': raise HTTPException(409,'Finish Cvent Microsoft SSO and reach an authenticated Cvent page before saving login')
    try: facts=asyncio.run(capture_login_facts())
    except Exception as e: raise HTTPException(500,f'Could not confirm Steel login cookies: {e}')
    if not facts.get('organization_id'): raise HTTPException(409,'Authenticated Cvent organization cookie was not found')
    data={**facts,'authenticated_at':now(),'cookie_store':str(DATA/'steel-profile-local'/'Default'/'Cookies')}
    atomic_json(AUTH_SETTINGS,data); os.chmod(AUTH_SETTINGS,0o600); append_log(f'Saved Cvent organization ID {facts["organization_id"]} and confirmed persistent Cvent/Microsoft SSO cookies')
    return {'ok':True,**auth_settings()}

@app.post('/api/start')
def start():
    ensure()
    with _lock:
        if running(): raise HTTPException(409,'A job is already running')
        if read_gate().get('ownership')!='AGENT': raise HTTPException(409,'Return browser control to the agent before starting')
        if not (CURRENT/'input.xlsx').exists(): raise HTTPException(400,'Upload the RR workbook first')
        try:load_scope_manifest()
        except Exception as e:raise HTTPException(500,f'Automation scope is invalid: {e}')
        locked_url=authorized_target_url();browser=chrome_status()
        if not browser.get('running'):browser=launch_chrome(locked_url or 'https://app.cvent.com/')
        if not browser.get('running'): raise HTTPException(500,browser.get('error','Chrome failed'))
        runtime=ensure_browser_runtime(full_probe=False);live=local_runtime_probe(runtime)
        if locked_url and event_key_from_url(live.get('url',''))!=event_key_from_url(locked_url):raise HTTPException(409,'The live browser is not on the authorized Cvent event; use OPEN BROWSER and verify it before starting')
        prompt=render_prompt(); pid=spawn_pi(prompt)
    return {'ok':True,'pid':pid,'browser':browser}

@app.post('/api/continue')
def continue_job():
    ensure()
    with _lock:
        if running(): raise HTTPException(409,'CVENT Agent is already running')
        if read_gate().get('ownership')!='AGENT': raise HTTPException(409,'Return browser control to the agent before continuing')
        try:load_scope_manifest()
        except Exception as e:raise HTTPException(500,f'Automation scope is invalid: {e}')
        st=read_json(STATE,{})
        locked_url=authorized_target_url()
        launch_chrome(locked_url or 'https://app.cvent.com/')
        runtime=ensure_browser_runtime(full_probe=True)
        auth=chrome_auth()
        if auth.get('auth_status')!='authenticated':
            where='Microsoft SSO/MFA' if auth.get('auth_status')=='microsoft_sso' else 'Cvent login'
            st.update({'status':'login_required','current_stage':'login','current_action':f'Finish {where} in OPEN BROWSER, including Stay signed in, before CONTINUE','updated_at':now()}); atomic_json(STATE,st)
            append_log(f'CONTINUE ignored: authentication incomplete at {where}')
            raise HTTPException(409,f'Authentication is still at {where}. Click OPEN BROWSER, finish Microsoft SSO/MFA and Stay signed in, then press CONTINUE.')
        msg='POLICY REPLACEMENT: discard any prior domain list or scope assumptions. The following complete prompt is now controlling and Intake Emerald is a fail-closed boundary.\n\n'+render_prompt()
        pid=spawn_pi(msg,resume=True)
    return {'ok':True,'pid':pid}

@app.post('/api/open-browser')
def open_browser():
    ensure(); gate=read_gate()
    # While Pi owns a live runtime, this button only reveals the viewer; it may
    # never navigate the shared page behind the agent's back.
    if running() and gate.get('ownership')=='AGENT':
        runtime=ensure_browser_runtime(full_probe=False)
        return {**chrome_status(),'browserRuntime':runtime,'displayOnly':True}
    result=launch_chrome(authorized_target_url() or 'https://app.cvent.com/')
    if result.get('running'):result['browserRuntime']=ensure_browser_runtime(full_probe=True)
    return result

@app.post('/api/browser/take-control')
def take_control():
    runtime=load_browser_runtime(BROWSER_RUNTIME_PATH); local_runtime_probe(runtime)
    request_user()
    with lock_file():
        pid=job_pid(); pids=process_tree(pid) if pid else []
        for process in pids:
            try:os.kill(process,signal.SIGSTOP)
            except ProcessLookupError:pass
        gate=read_gate(); gate.update({'ownership':'USER','desiredOwnership':'USER','activeActor':'USER','automationOwner':'USER','agentPaused':bool(pids),'pausedPids':pids,'transition':None,'browserRuntimeId':runtime['browserRuntimeId']});write_gate(gate)
    append_log('Human takeover enabled at a safe browser action boundary')
    return {'ok':True,'gate':read_gate()}

@app.post('/api/browser/return-to-agent')
def return_to_agent():
    runtime=load_browser_runtime(BROWSER_RUNTIME_PATH);shield_agent()
    with lock_file():
        try:
            viewer=local_runtime_probe(runtime)
            ego=tool_probe(runtime,['node','ego_direct.mjs','--runtime',str(BROWSER_RUNTIME_PATH),'--operation','snapshotText','--params','{}'])
            ego_page=tool_probe(runtime,['node','ego_direct.mjs','--runtime',str(BROWSER_RUNTIME_PATH),'--operation','pageInfo','--params','{}'])
            if not ego.get('ok') or not ego_page.get('ok'):raise RuntimeError('Fresh Ego browser read failed')
            lock=read_json(CURRENT/'authorized-target.json',{})
            if lock:
                expected_key=event_key_from_url(lock.get('url',''));viewer_key=event_key_from_url(viewer.get('url',''));ego_key=event_key_from_url((ego_page.get('page') or {}).get('url',''))
                if not expected_key or lock.get('event_key')!=expected_key or viewer_key!=expected_key or ego_key!=expected_key or AUTHORIZED_EVENT_NAME.lower() not in json.dumps(ego).lower():raise RuntimeError('Human left the authorized Cvent event; CVENT Agent remains paused')
            handoff={'browserRuntimeId':runtime['browserRuntimeId'],'viewer':viewer,'ego':ego,'egoPage':ego_page,'inspectedAt':now()};atomic_json(CURRENT/'human-handoff-state.json',handoff)
            gate=read_gate();paused=gate.get('pausedPids',[]);gate.update({'ownership':'AGENT','desiredOwnership':'AGENT','activeActor':'NONE','automationOwner':'PI_EGO','agentPaused':False,'pausedPids':[],'transition':None});write_gate(gate)
            for process in reversed(paused):
                try:os.kill(process,signal.SIGCONT)
                except ProcessLookupError:pass
        except Exception:
            gate=read_gate();gate.update({'ownership':'NONE','desiredOwnership':'AGENT','activeActor':'NONE','automationOwner':'NONE','transition':'RETURN_BLOCKED','agentPaused':True});write_gate(gate);raise
    append_log(f'Returned browser to CVENT Agent after fresh Ego verification: {viewer["title"]}')
    return {'ok':True,'gate':read_gate(),'state':handoff}

def event_key_from_url(url):
    from urllib.parse import parse_qs
    q={key.lower():value for key,value in parse_qs(urlparse(url).query).items()}
    for key in ('evtstub','eventid','event'):
        if q.get(key):return q[key][0].lower()
    match=re.search(r'/events/([0-9a-f-]{20,})',urlparse(url).path,re.I)
    return match.group(1).lower() if match else None

@app.post('/api/stop-agent')
def stop_agent():
    global _proc
    with _lock:
        pid=job_pid()
        if pid:
            stop_process_tree(pid); _proc=None
            append_log(f'CVENT Agent process tree {pid} stopped by operator')
        steel=steel_command('release',timeout=60)
        initialize_gate()
        try:BROWSER_RUNTIME_PATH.unlink()
        except FileNotFoundError:pass
        st=read_json(STATE,{})
        st.update({'status':'stopped','current_stage':'stopped','current_action':'Build and Steel browser stopped','pi_pid':None,'process_started_at':None,'updated_at':now()}); atomic_json(STATE,st)
        append_log('Steel OSS browser stopped; persistent profile preserved')
    return {'ok':True,'steel':steel}
